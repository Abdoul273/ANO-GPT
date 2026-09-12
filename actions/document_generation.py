"""Rédaction de documents complets par le modèle Azure « document ».

Le texte est écrit par le spécialiste Azure configuré, puis rendu dans le
format demandé.  Markdown, texte et HTML sont toujours disponibles ; DOCX, PDF
et PPTX ne le sont que si la bibliothèque correspondante est installée — dans
ce cas on rend un Markdown et on le dit clairement plutôt que d'échouer.
"""
from __future__ import annotations

import re
import time
from pathlib import Path

from core.azure_specialists import text

from core import action_kit as kit

OUTPUT_DIR = Path.home() / "Documents" / "ANO-GPT"

FORMATS = {"md", "txt", "html", "docx", "pdf", "pptx"}

_SYSTEM = (
    "Tu es un rédacteur professionnel. Produis un document complet, structuré "
    "et prêt à être livré, en Markdown (titres #, listes, tableaux si utile). "
    "N'ajoute aucun commentaire sur ton propre travail, aucune balise de code "
    "autour du document. Réponds dans la langue de la demande, français par défaut."
)


def _slug(value: str) -> str:
    return re.sub(r"[^\w-]+", "_", value, flags=re.UNICODE).strip("_")[:48] or "document"


def _strip_fence(body: str) -> str:
    body = body.strip()
    if body.startswith("```"):
        body = re.sub(r"^```[a-zA-Z]*\n", "", body)
        body = re.sub(r"\n```$", "", body)
    return body.strip()


def _write_html(body: str, path: Path, title: str) -> None:
    lines = []
    for raw in body.splitlines():
        line = raw.rstrip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            lines.append(f"<h{level}>{heading.group(2)}</h{level}>")
        elif re.match(r"^[-*]\s+", line):
            lines.append(f"<li>{line[2:]}</li>")
        elif line:
            lines.append(f"<p>{line}</p>")
    path.write_text(
        "<!doctype html><meta charset='utf-8'>"
        f"<title>{title}</title>"
        "<style>body{max-width:46em;margin:3em auto;font:16px/1.6 system-ui;padding:0 1em}</style>"
        + "\n".join(lines),
        encoding="utf-8",
    )


def _write_docx(body: str, path: Path) -> bool:
    try:
        from docx import Document  # type: ignore
    except Exception:
        return False
    document = Document()
    for raw in body.splitlines():
        line = raw.rstrip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
        elif re.match(r"^[-*]\s+", line):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line:
            document.add_paragraph(line)
    document.save(str(path))
    return True


def _write_pdf(body: str, path: Path) -> bool:
    try:
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.styles import getSampleStyleSheet  # type: ignore
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # type: ignore
    except Exception:
        return False
    styles = getSampleStyleSheet()
    flow = []
    for raw in body.splitlines():
        line = raw.rstrip().replace("&", "&amp;").replace("<", "&lt;")
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            flow.append(Paragraph(heading.group(2), styles[f"Heading{min(len(heading.group(1)), 4)}"]))
        elif line:
            flow.append(Paragraph(line, styles["BodyText"]))
        else:
            flow.append(Spacer(1, 8))
    SimpleDocTemplate(str(path), pagesize=A4).build(flow)
    return True


def _write_pptx(body: str, path: Path, title: str) -> bool:
    try:
        from pptx import Presentation  # type: ignore
    except Exception:
        return False
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = title
    current = None
    for raw in body.splitlines():
        line = raw.rstrip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            current = deck.slides.add_slide(deck.slide_layouts[1])
            current.shapes.title.text = heading.group(2)
        elif line and current is not None:
            frame = current.placeholders[1].text_frame
            paragraph = frame.paragraphs[0] if not frame.text else frame.add_paragraph()
            paragraph.text = re.sub(r"^[-*]\s+", "", line)
    deck.save(str(path))
    return True


@kit.action("generate_document")
def generate_document(parameters: dict | None = None, player=None) -> str:
    parameters = parameters or {}
    subject = " ".join(str(parameters.get("subject") or parameters.get("prompt") or "").split())
    if not subject:
        return "Précisez le sujet du document à rédiger."
    fmt = str(parameters.get("format") or "md").strip().lower().lstrip(".")
    if fmt in {"markdown"}:
        fmt = "md"
    if fmt in {"text"}:
        fmt = "txt"
    if fmt in {"word"}:
        fmt = "docx"
    if fmt in {"powerpoint", "ppt", "slides"}:
        fmt = "pptx"
    if fmt not in FORMATS:
        fmt = "md"
    title = " ".join(str(parameters.get("title") or subject).split())[:80]

    instructions = str(parameters.get("instructions") or "").strip()
    prompt = f"Rédige un document intitulé « {title} ».\n\nSujet : {subject}"
    if instructions:
        prompt += f"\n\nContraintes : {instructions}"
    if fmt == "pptx":
        prompt += (
            "\n\nStructure-le en diapositives : un titre de niveau ## par diapositive, "
            "puis 3 à 5 puces courtes."
        )
    try:
        body = _strip_fence(text("document", prompt, system=_SYSTEM, timeout=180))
    except Exception as exc:
        return f"Rédaction Azure échouée : {exc}"
    if not body:
        return "Le modèle Azure n'a produit aucun texte."

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stem = f"{_slug(title)}_{int(time.time())}"
    path = OUTPUT_DIR / f"{stem}.{fmt}"
    note = ""
    try:
        if fmt in {"md", "txt"}:
            path.write_text(body, encoding="utf-8")
        elif fmt == "html":
            _write_html(body, path, title)
        elif fmt == "docx" and not _write_docx(body, path):
            path = OUTPUT_DIR / f"{stem}.md"
            path.write_text(body, encoding="utf-8")
            note = " (DOCX indisponible : installez python-docx ; document rendu en Markdown)"
        elif fmt == "pdf" and not _write_pdf(body, path):
            path = OUTPUT_DIR / f"{stem}.md"
            path.write_text(body, encoding="utf-8")
            note = " (PDF indisponible : installez reportlab ; document rendu en Markdown)"
        elif fmt == "pptx" and not _write_pptx(body, path, title):
            path = OUTPUT_DIR / f"{stem}.md"
            path.write_text(body, encoding="utf-8")
            note = " (PPTX indisponible : installez python-pptx ; document rendu en Markdown)"
    except Exception as exc:
        return f"Écriture du document impossible : {exc}"

    if player is not None and hasattr(player, "show_generated_artifact_preview"):
        try:
            player.show_generated_artifact_preview("document", title[:80] or "Document créé", str(path))
        except Exception:
            pass
    if parameters.get("open_after"):
        try:
            kit.spawn(['xdg-open', str(path)])
        except Exception:
            pass
    return f"Document « {title} » enregistré dans : {path}{note}"
